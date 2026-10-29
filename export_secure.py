import getpass
import io
import os
import sys

import msoffcrypto
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.protection import WorkbookProtection
from werkzeug.security import check_password_hash

import app as a

d = a.load()

# ---- 1. only the admin can run this ----
print("SECURE EXPORT: only the admin can run this.")
name = input("Admin username: ")
pass_in = getpass.getpass("Admin password: ")
adm = a.find_user(d, name)
if not adm or adm["role"] != "admin" or not check_password_hash(adm["password"], pass_in):
    sys.exit("Access denied.")

# ---- 2. password that locks the Excel file ----
pw = getpass.getpass("Choose a password for the Excel file (min 8 characters): ")
if len(pw) < 8 or pw != getpass.getpass("Type it again: "):
    sys.exit("Passwords did not match or are shorter than 8 characters.")

students = {s["roll"]: s for s in d["students"]}
subjects = {s["code"]: s for s in d["subjects"]}
wb = Workbook()
HEAD = PatternFill("solid", fgColor="1F4E78")
FILL = {"RED": "F8B4B4", "YELLOW": "FFF2A8", "GREEN": "B7F0D2", "DEFAULT": "FFF2A8"}


def sheet(title, headers, rows, first=False, status_col=None):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    if status_col:
        for row in ws.iter_rows(min_row=2, min_col=status_col, max_col=status_col):
            v = str(row[0].value or "").split(" ")[0].upper().strip(":")
            if v in FILL:
                row[0].fill = PatternFill("solid", fgColor=FILL[v])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.columns, 1):
        w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 3, 12), 45)


def state(user, default_pw, who):
    if check_password_hash(user["password"], default_pw):
        return "Default (same as " + ("roll no" if who == "student" else "username") + ")"
    return "Changed or custom"


# ---- 3. login sheets (first three) ----
rows = []
for st in d["students"]:
    u = a.find_user(d, st["roll"])
    rows.append([st["roll"], st["name"], st["klass"], state(u, u["username"], "student") if u else "No login"])
sheet("Student Logins", ["Username (Roll No)", "Name", "Class", "Password Status"], rows, first=True)

rows = []
for u in d["users"]:
    if u["role"] == "teacher":
        subs = ", ".join(s["code"] for s in d["subjects"] if s["teacher"] == u["username"])
        rows.append([u["username"], u["name"], subs, state(u, u["username"], "teacher")])
sheet("Teacher Logins", ["Username", "Name", "Subjects", "Password Status"], rows)

rows = []
for u in d["users"]:
    if u["role"] == "admin":
        st = "DEFAULT admin123: CHANGE IT NOW" if check_password_hash(u["password"], "admin123") else "Changed"
        rows.append([u["username"], u["name"], st])
sheet("Admin Login", ["Username", "Name", "Password Status"], rows, status_col=3)

# ---- 4. data sheets ----
rows = []
for st in d["students"]:
    _, o = a.student_stats(d, st)
    rows.append([st["roll"], st["name"], st["klass"], o["percent"] if o["total"] else "",
                 o["present"], o["total"], o["status"].upper() if o["total"] else "NO DATA"])
sheet("Students", ["Roll No", "Name", "Class", "Attendance %", "Present", "Total Classes", "Status"], rows, status_col=7)

tn = {u["username"]: u["name"] for u in d["users"] if u["role"] == "teacher"}
sheet("Subjects", ["Code", "Subject", "Class", "Teacher Username", "Teacher Name"],
      [[s["code"], s["name"], s["klass"], s["teacher"], tn.get(s["teacher"], "")] for s in d["subjects"]])

rows = []
for r in sorted(d["attendance"], key=lambda x: (x["date"], x["subject"])):
    for roll, present in r["status"].items():
        rows.append([r["date"], r["subject"], subjects.get(r["subject"], {}).get("name", ""),
                     roll, students.get(roll, {}).get("name", ""), "Present" if present else "Absent"])
sheet("Attendance", ["Date", "Subject Code", "Subject", "Roll No", "Name", "Status"], rows)

rows = []
for st in d["students"]:
    subs, _ = a.student_stats(d, st)
    for x in subs:
        rows.append([st["roll"], st["name"], st["klass"], x["name"], x["present"], x["total"],
                     x["percent"] if x["total"] else "", x["status"].upper() if x["total"] else "NO DATA"])
sheet("Attendance Summary", ["Roll No", "Name", "Class", "Subject", "Present", "Total", "Attendance %", "Status"],
      rows, status_col=8)

rows = []
for key, m in d["marks"].items():
    code, roll = key.split("|")
    st = students.get(roll, {})
    rows.append([roll, st.get("name", ""), st.get("klass", ""), subjects.get(code, {}).get("name", code),
                 m.get("mid1", ""), m.get("mid2", ""), m.get("internal", ""), m.get("test", ""), m.get("sem", "")])
sheet("Marks", ["Roll No", "Name", "Class", "Subject", "Mid 1", "Mid 2", "Internal", "Class Test", "Semester"], rows)

# ---- 5. lock every sheet (read-only, filter and sort still allowed) ----
for ws in wb.worksheets:
    ws.protection.sheet = True
    ws.protection.set_password(pw)
    ws.protection.autoFilter = False
    ws.protection.sort = False
wb.security = WorkbookProtection(lockStructure=True)
wb.security.set_workbook_password(pw)

# ---- 6. encrypt the file (AES): never written unencrypted to disk ----
buf = io.BytesIO()
wb.save(buf)
buf.seek(0)
os.makedirs("exports", exist_ok=True)
path = os.path.join("exports", "secure_database.xlsx")
try:
    with open(path, "wb") as out:
        msoffcrypto.OfficeFile(buf).encrypt(pw, out)
except PermissionError:
    sys.exit("Close secure_database.xlsx in Excel first, then run this again.")

print("\nSaved (encrypted):", os.path.abspath(path))
print("Excel will ask for the file password when it opens.")
try:
    os.startfile(path)
except Exception:
    pass
